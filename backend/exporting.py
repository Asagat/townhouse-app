# backend/exporting.py
"""Экспорт списков (регистров и документов) в Excel/CSV (роадмап 2.2).

Экспорт встроен в уже существующий универсальный список `GET /api/{resource}`:
клиент передаёт те же параметры фильтрации/сортировки, что и при просмотре, плюс
два служебных:

    _export=xlsx|csv      — вместо JSON вернуть файл;
    _columns=<JSON>       — какие колонки выгружать и как они называются
                            (список вида [{"key":"amount","label":"Сумма"}, …]).

Колонки присылает фронтенд — ровно те, что пользователь видит (с учётом скрытия/
порядка колонок), поэтому выгрузка совпадает с таблицей. Значения берутся из того же
сериализатора, что и список: ключ колонки — путь по вложенным полям (`account.account_number`).

Правила нормализации значений:
  - None/пусто → пустая ячейка;
  - boolean → «Да»/«Нет»;
  - Decimal → число (в Excel — как число);
  - объекты-справочники (`{"id":..,"name":..}`) → человекочитаемый текст без id;
  - списки → значения через запятую.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from urllib.parse import quote

from fastapi import HTTPException
from fastapi.responses import Response

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Разделитель CSV «;» — привычный для Excel в русской локали (иначе весь ряд
# может попасть в одну колонку). Кодировка — UTF-8 с BOM (см. _to_csv).
CSV_DELIMITER = ";"

# Подписи для кодовых значений (там, где список показывает русский текст через
# собственный format): экспорт должен читаться так же, как таблица.
_FIELD_VALUE_LABELS: dict[str, dict[str, str]] = {
    "doc_kind": {"monthly": "Регулярные", "oneoff": "Персональные"},
    "status": {"active": "Действующий", "archived": "Архивный"},
}


def _label_for(key: str, value: Any) -> Any:
    mapping = _FIELD_VALUE_LABELS.get(key)
    if mapping and isinstance(value, str):
        return mapping.get(value, value)
    return value


def get_by_path(record: dict, path: str) -> Any:
    """Значение по пути `a.b.c` во вложенном словаре (сериализованная строка)."""
    current: Any = record
    for part in path.split("."):
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return None
        if current is None:
            return None
    return current


def _dict_to_text(value: dict) -> str:
    """Человекочитаемый текст из объекта-справочника (без служебного id)."""
    for key in ("name", "full_name", "services_type", "account_number", "title"):
        text = value.get(key)
        if text not in (None, ""):
            return str(text)
    parts = [str(v) for k, v in value.items() if k != "id" and v not in (None, "")]
    return " / ".join(parts)


def normalize_value(value: Any) -> Any:
    """Приводит значение к «ячейке» (см. правила в шапке модуля)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return _dict_to_text(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(normalize_value(v) or "") for v in value)
    return str(value)


def resolve_columns(rows: list[dict], raw_columns: str | None) -> list[tuple[str, str]]:
    """Список колонок (ключ, заголовок): из _columns или запасной вариант по ключам строки."""
    if raw_columns:
        try:
            parsed = json.loads(raw_columns)
        except (json.JSONDecodeError, TypeError):
            parsed = None
        if isinstance(parsed, list):
            columns: list[tuple[str, str]] = []
            for item in parsed:
                if isinstance(item, dict) and item.get("key"):
                    key = str(item["key"])
                    label = str(item.get("label") or key)
                    columns.append((key, label))
            if columns:
                return columns
    # Запасной вариант (клиент не прислал колонок): ключи первой строки.
    if rows:
        return [(str(k), str(k)) for k in rows[0].keys()]
    return []


def _build_table(rows: list[dict], columns: list[tuple[str, str]]) -> tuple[list[str], list[list[Any]]]:
    header = [label for _, label in columns]
    body = [
        [normalize_value(_label_for(key, get_by_path(row, key))) for key, _ in columns]
        for row in rows
    ]
    return header, body


def _to_csv(header: list[str], body: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=CSV_DELIMITER, lineterminator="\r\n")
    writer.writerow(header)
    for row in body:
        writer.writerow(["" if cell is None else cell for cell in row])
    # BOM (utf-8-sig) — чтобы Excel открывал кириллицу без «кракозябр».
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _to_xlsx(header: list[str], body: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Экспорт"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in body:
        ws.append(row)
    # Ширины колонок — по самому длинному значению (в разумных пределах).
    for idx, title in enumerate(header, start=1):
        longest = len(str(title))
        for row in body:
            value = row[idx - 1]
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 8), 60)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_export_response(
    resource: str, rows: list[dict], raw_columns: str | None, fmt: str
) -> Response:
    """Формирует файл экспорта из уже сериализованных строк списка."""
    fmt = (fmt or "").strip().lower()
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(status_code=422, detail="Формат экспорта должен быть csv или xlsx")

    columns = resolve_columns(rows, raw_columns)
    header, body = _build_table(rows, columns)

    if fmt == "csv":
        content = _to_csv(header, body)
        media_type = CSV_MEDIA_TYPE
        ext = "csv"
    else:
        content = _to_xlsx(header, body)
        media_type = XLSX_MEDIA_TYPE
        ext = "xlsx"

    filename = f"{resource}_{datetime.now():%Y%m%d_%H%M%S}.{ext}"
    # filename* (RFC 5987) — корректная кириллица/спецсимволы в имени при скачивании.
    disposition = (
        f'attachment; filename="{resource}.{ext}"; '
        f"filename*=UTF-8''{quote(filename)}"
    )
    return Response(content=content, media_type=media_type, headers={"Content-Disposition": disposition})
