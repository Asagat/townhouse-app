# tests/test_registers_settlements_vs_cash.py
"""Взаиморасчёты (accounts_register) ↔ жительские взносы (Касса-Приход файла).

В старой базе жительский взнос сразу погашал долг; в новой один взнос даёт
cash_register.income (касса) и accounts_register.expense (списание долга). Сверяем
«расход Взаиморасчётов» с файловым «жительским приходом» (нетто, сторно вычитается)
на срезах:
  1) итог;
  2) по квартире;
  3) по месяцу (дата операции в accounts == месяц платежа кассы);
  4) месяц × квартира.

Переплата в accounts_register expense не попадает, но на текущем наборе переплат нет
(все 17 л/с в долг), поэтому строгие равенства корректны.
"""

import os
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import text

_BACKEND_DIR = Path(__file__).resolve().parents[2]
_REPO_ROOT = Path(__file__).resolve().parents[3]

try:
    import openpyxl  # noqa: F401
except Exception:  # pragma: no cover
    openpyxl = None  # type: ignore


def _source_path() -> Path | None:
    env = os.getenv("MIGRATION_SRC_XLSX")
    if env:
        p = Path(env)
        return p if p.exists() else None
    p = _REPO_ROOT / "templates" / "Миграция данных FTH.xlsx"
    return p if p.exists() else None


def _ym(d) -> str:
    return f"{d.year:04d}-{d.month:02d}"


@pytest.fixture
def settle(db):
    if openpyxl is None:
        pytest.skip("Нет openpyxl — сверка взаиморасчётов недоступна")
    src = _source_path()
    if src is None:
        pytest.skip("Файл-источник не найден (задайте MIGRATION_SRC_XLSX)")
    has = db.execute(text("SELECT EXISTS (SELECT 1 FROM accounts_register)")).scalar()
    if not has:
        pytest.skip("БД без взаиморасчётов")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    return wb, db


def _source_resident_income(wb):
    """:(месяц,кв) -> приход жителей нетто. Возврат dict и карту кода квартиры."""
    kv_code = {}
    for r in wb["Квартиры"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        kv_code[int(r[0])] = r[2]
    inc = defaultdict(Decimal)  # (ym, kv)
    for r in wb["Касса-Приход"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        if not hasattr(r[6], "year"):
            continue
        s = str(r[1]).strip() if r[1] is not None else ""
        if s in ("", "None", "0", "0.0"):
            continue
        try:
            key = str(r[1]).split(".")[0] if isinstance(r[1], float) else s
            kv = kv_code.get(int(float(key)))
        except (TypeError, ValueError):
            continue
        if kv is None:
            continue
        amt = Decimal(str(r[5]))
        cm = str(r[8] or "")
        storno = amt < 0 or "[СТОРНО]" in cm or "сторно" in cm.lower()
        inc[(_ym(r[6]), int(kv))] += (-abs(amt) if storno else abs(amt))
    return inc, kv_code


def _db_expense(db):
    exp = defaultdict(Decimal)
    for mm, kv, v in db.execute(text(
        "SELECT to_char(ar.operation_date,'YYYY-MM'), ap.apartment_number, ar.expense "
        "FROM accounts_register ar "
        "JOIN accounts a ON a.id = ar.account_id "
        "JOIN apartments ap ON ap.id = a.apartment_id")).fetchall():
        exp[(mm, int(kv))] += Decimal(str(v))
    return exp


def test_total(db, settle):
    inc, _ = _source_resident_income(settle[0])
    exp = _db_expense(db)
    a = sum(inc.values()); b = sum(exp.values())
    _eq(a, b, "Итог (взаиморасчёты ↔ жительский приход)")


def test_per_apartment(db, settle):
    inc, _ = _source_resident_income(settle[0])
    exp = _db_expense(db)
    a = {kv: sum(v for (_, k), v in inc.items() if k == kv) for kv in range(1, 18)}
    b = {kv: sum(v for (_, k), v in exp.items() if k == kv) for kv in range(1, 18)}
    _diff_dict(a, b, "По квартире")


def test_per_month(db, settle):
    inc, _ = _source_resident_income(settle[0])
    exp = _db_expense(db)
    a = {m: sum(v for (mm, _), v in inc.items() if mm == m)
         for m in sorted(set(k[0] for k in inc))}
    b = {m: sum(v for (mm, _), v in exp.items() if mm == m)
         for m in sorted(set(k[0] for k in exp) | set(k[0] for k in inc))}
    _diff_dict(a, b, "По месяцу")


def test_per_month_per_apartment(db, settle):
    inc, _ = _source_resident_income(settle[0])
    exp = _db_expense(db)
    keys = sorted(set(inc) | set(exp))
    diffs = [
        (k, inc.get(k, Decimal(0)), exp.get(k, Decimal(0)))
        for k in keys if abs(inc.get(k, Decimal(0)) - exp.get(k, Decimal(0))) >= Decimal("0.01")
    ]
    assert not diffs, (
        "Месяц × квартира: расхождений "
        f"{len(diffs)}:\n" + "\n".join(f"  {k}: приход {a:,.2f} vs списание {b:,.2f}" for k, a, b in diffs[:20])
    )


def _eq(a: Decimal, b: Decimal, label: str):
    assert abs(a - b) < Decimal("0.01"), f"{label}: приход {a:,.2f} vs списание {b:,.2f}"


def _diff_dict(a, b, label):
    keys = sorted(set(a) | set(b))
    diffs = [(k, a.get(k, Decimal(0)), b.get(k, Decimal(0))) for k in keys
             if abs(a.get(k, Decimal(0)) - b.get(k, Decimal(0))) >= Decimal("0.01")]
    assert not diffs, (
        f"{label}: расхождений {len(diffs)}:\n" +
        "\n".join(f"  {k}: приход {x:,.2f} vs списание {y:,.2f} (Δ={y - x:,.2f})"
                  for k, x, y in diffs[:20])
    )
