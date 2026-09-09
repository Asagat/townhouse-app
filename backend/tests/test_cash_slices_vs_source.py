# tests/test_cash_slices_vs_source.py
"""Денежный регистр (cash_register): месячные срезы файл-источника ↔ БД.

Сверяет листы «Касса-Приход»/«Касса-Расход» с cash_register по месяцу операции
(operation_date в БД = дата операции транзакции):
  1) приход по месяцу  (месяц × сумма прихода, нетто: сторно вычитается);
  2) расход по месяцу  (месяц × сумма расхода);
  3) помесячный приход по квартире (жительские взносы, нетто сторно).

Начальный остаток кассы («Начальный остаток …») участвует в приходе на своём месяце
как обычная операция. Если файла-источника / openpyxl нет — тест пропускается.
"""

import importlib.util
import os
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import text

try:
    import openpyxl  # noqa: F401
except Exception:  # pragma: no cover
    openpyxl = None  # type: ignore

_BACKEND_DIR = Path(__file__).resolve().parents[1]
_REPO_ROOT = Path(__file__).resolve().parents[2]


def _source_path() -> Path | None:
    env = os.getenv("MIGRATION_SRC_XLSX")
    if env:
        p = Path(env)
        return p if p.exists() else None
    p = _REPO_ROOT / "templates" / "Миграция данных FTH.xlsx"
    return p if p.exists() else None


def _ym(d) -> str:
    return f"{d.year:04d}-{d.month:02d}"


@pytest.fixture
def checker(db):
    if openpyxl is None:
        pytest.skip("Нет пакета openpyxl — сверка денежного регистра недоступна")
    src = _source_path()
    if src is None:
        pytest.skip("Файл-источник не найден (задайте MIGRATION_SRC_XLSX)")
    has = db.execute(text("SELECT EXISTS (SELECT 1 FROM cash_register)")).scalar()
    if not has:
        pytest.skip("БД без кассы — сверка не применима")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    return wb, db


def _src_needed(wb):
    return None




def test_cash_monthly_income_matches(db, checker):
    """Приход по месяцам (нетто): файл == БД."""
    wb, _ = checker
    src = defaultdict(Decimal)
    for r in wb["Касса-Приход"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        if not hasattr(r[6], "year"):
            continue
        amt = Decimal(str(r[5]))
        cm = str(r[8] or "")
        storno = amt < 0 or "[СТОРНО]" in cm or "сторно" in cm.lower()
        src[_ym(r[6])] += (-abs(amt) if storno else abs(amt))
    db_rows = defaultdict(Decimal)
    for mm, v in db.execute(text(
        "SELECT to_char(operation_date,'YYYY-MM'), COALESCE(SUM(income),0) "
        "FROM cash_register GROUP BY 1")).fetchall():
        db_rows[mm] += Decimal(str(v))
    _assert_empty_diff(src, db_rows, "Приход по месяцам")


def test_cash_monthly_expense_matches(db, checker):
    """Расход по месяцам: файл == БД."""
    wb, _ = checker
    src = defaultdict(Decimal)
    for r in wb["Касса-Расход"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        if not hasattr(r[7], "year"):
            continue
        src[_ym(r[7])] += Decimal(str(r[6]))  # Сумма (индекс 6)
    db_rows = defaultdict(Decimal)
    for mm, v in db.execute(text(
        "SELECT to_char(operation_date,'YYYY-MM'), COALESCE(SUM(expense),0) "
        "FROM cash_register GROUP BY 1")).fetchall():
        db_rows[mm] += Decimal(str(v))
    _assert_empty_diff(src, db_rows, "Расход по месяцам")


def test_cash_income_per_apartment_per_month(db, checker):
    """Приход по квартире и месяцу (нетто, сторно вычитается): файл == БД."""
    wb, _ = checker
    kv_code = {}
    for r in wb["Квартиры"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        kv_code[int(r[0])] = r[2]
    src = defaultdict(Decimal)  # (ym, kv_num)
    for r in wb["Касса-Приход"].iter_rows(values_only=True):
        if r[0] is None or str(r[0]).startswith("Код"):
            continue
        if not hasattr(r[6], "year"):
            continue
        code = r[1]
        s = str(code).strip() if code is not None else ""
        if s in ("", "None", "0", "0.0"):
            continue  # операция без квартиры — не жительская
        try:
            key = str(code).split(".")[0] if isinstance(code, float) else s.split(".")[0]
            kv = kv_code.get(int(key))
        except (TypeError, ValueError):
            kv = None
        if kv is None:
            continue
        amt = Decimal(str(r[5]))
        cm = str(r[8] or "")
        storno = amt < 0 or "[СТОРНО]" in cm or "сторно" in cm.lower()
        src[(_ym(r[6]), int(kv))] += (-abs(amt) if storno else abs(amt))
    db_rows = defaultdict(Decimal)
    for mm, kvnum, v in db.execute(text(
        "SELECT to_char(c.operation_date,'YYYY-MM'), ap.apartment_number, COALESCE(SUM(c.income),0) "
        "FROM cash_register c "
        "JOIN accounts a ON a.id=c.account_id JOIN apartments ap ON ap.id=a.apartment_id "
        "GROUP BY 1, 2")).fetchall():
        db_rows[(mm, int(kvnum))] += Decimal(str(v))
    _assert_empty_diff(src, db_rows, "Приход по квартире/месяцу")


def _assert_empty_diff(src, db_rows, label):
    keys = sorted(set(src) | set(db_rows))
    diffs = [
        (k, src.get(k, Decimal(0)), db_rows.get(k, Decimal(0)))
        for k in keys
        if abs(src.get(k, Decimal(0)) - db_rows.get(k, Decimal(0))) >= Decimal("0.01")
    ]
    assert not diffs, (
        f"{label}: расхождений {len(diffs)}:\n"
        + "\n".join(f"  {k}: файл {a:,.2f} vs БД {b:,.2f} (Δ={b - a:,.2f})"
                    for k, a, b in diffs[:20])
    )
