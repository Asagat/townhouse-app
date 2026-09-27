# backend/tests/test_export.py
"""Экспорт списков (регистров и документов) в Excel/CSV (роадмап 2.2).

Проверяем встроенный в универсальный список эндпоинт GET /api/{resource}?_export=…:
  - возвращается файл (application/…), а не JSON;
  - колонки берутся из _columns (ключ + заголовок), значения — из сериализатора;
  - применённые фильтры учитываются (выгружается отфильтрованное подмножество);
  - доступ проверяется тем же require_resource_access, что и список.
"""

import io
import json
from datetime import datetime
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from auth import create_access_token
from exporting import build_export_response
from models import Transaction, TransactionTypeEnum, UserRole


def _headers(admin):
    return {"Authorization": f"Bearer {create_access_token(admin)}"}


def _columns(*pairs) -> str:
    return json.dumps([{"key": k, "label": v} for k, v in pairs])


def test_export_owners_csv_with_columns_and_filter(db, user_factory, account_factory):
    admin = user_factory("expadmin", UserRole.admin)
    account_factory("bexp1")
    account_factory("bexp2")
    client = TestClient(app)

    resp = client.get(
        "/api/owners",
        params={
            "_export": "csv",
            "_columns": _columns(("full_name", "ФИО"), ("phone", "Телефон")),
            "full_name_like": "bexp",
        },
        headers=_headers(admin),
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]

    text = resp.content.decode("utf-8-sig")
    lines = [ln for ln in text.split("\r\n") if ln]
    # Заголовок — из переданных подписей колонок.
    assert lines[0] == "ФИО;Телефон"
    # Только отфильтрованные записи (2 созданных «bexpN T»), без лишних.
    assert len(lines) == 3
    assert all("bexp" in ln for ln in lines[1:])


def test_export_filter_narrows_rows(db, user_factory, account_factory):
    """Фильтр из UI (like) реально сужает выгрузку."""
    admin = user_factory("expfil", UserRole.admin)
    account_factory("bfx1")
    account_factory("bfx2")
    client = TestClient(app)

    resp = client.get(
        "/api/owners",
        params={"_export": "csv", "_columns": _columns(("full_name", "ФИО")), "full_name_like": "bfx1"},
        headers=_headers(admin),
    )
    assert resp.status_code == 200
    lines = [ln for ln in resp.content.decode("utf-8-sig").split("\r\n") if ln]
    assert lines == ["ФИО", "bfx1 T"]


def test_export_transactions_xlsx(db, user_factory, account_factory):
    admin = user_factory("expxls", UserRole.admin)
    acc = account_factory("bxls1")
    db.add(Transaction(
        account_id=acc["account_id"],
        cash_point_id=acc["cash_point_id"],
        transaction_type=TransactionTypeEnum.in_cash,
        amount=Decimal("1234.50"),
        transaction_date=datetime(2026, 9, 1, 10, 0),
    ))
    db.commit()
    client = TestClient(app)

    resp = client.get(
        "/api/transactions",
        params={
            "_export": "xlsx",
            "_columns": _columns(("amount", "Сумма"), ("transaction_type", "Тип операции")),
            "amount_gte": "1000",
        },
        headers=_headers(admin),
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Сумма"
    assert ws.cell(row=1, column=2).value == "Тип операции"
    # Числовая сумма остаётся числом (а не строкой) — удобно считать в Excel.
    assert ws.cell(row=2, column=1).value == 1234.5


def test_export_bad_format_rejected(db, user_factory):
    admin = user_factory("expbad", UserRole.admin)
    client = TestClient(app)
    resp = client.get("/api/owners", params={"_export": "pdf"}, headers=_headers(admin))
    assert resp.status_code == 422


def test_export_requires_authorization(db):
    client = TestClient(app)
    resp = client.get("/api/owners", params={"_export": "csv"})
    assert resp.status_code in (401, 403)


def test_export_value_normalization():
    """Вложенные пути, кодовые подписи и булево — без обращения к БД (чистая функция)."""
    rows = [{
        "doc_kind": "oneoff",
        "apartment": {"apartment_number": 5},
        "amount": 10.5,
        "is_active": True,
        "empty": None,
    }]
    cols = _columns(
        ("doc_kind", "Тип"),
        ("apartment.apartment_number", "№ кв"),
        ("amount", "Сумма"),
        ("is_active", "Активен"),
        ("empty", "Пусто"),
    )
    resp = build_export_response("accrual_documents", rows, cols, "csv")
    lines = resp.body.decode("utf-8-sig").split("\r\n")
    assert lines[0] == "Тип;№ кв;Сумма;Активен;Пусто"
    assert lines[1] == "Персональные;5;10.5;Да;"
